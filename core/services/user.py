import asyncio
import io
import json
import logging
import os
import secrets
import uuid
from datetime import UTC, date, datetime
from types import SimpleNamespace

import jwt
from dotenv import load_dotenv
from fastapi import HTTPException, Response, status
from fpdf import FPDF
from openpyxl import Workbook
from redis.asyncio import Redis
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from core.messages import USER_MESSAGES
from core.models.habit import ActivityLog, Habit
from core.models.routine import Routine
from core.models.user import Session, User
from core.schemas.user import UserResponse
from core.security import (
    create_token,
    decode_token,
    hash_password,
    hash_token,
    verify_password,
)
from core.utils import attach_activity, send_email_otp, send_welcome_mail
from db.database import AsyncSessionLocal

load_dotenv()

logger = logging.getLogger(__name__)

SIGNUP_OTP_EXPIRY_SECONDS = 600
RESET_PASSWORD_OTP_EXPIRY_SECONDS = 600

ENVIRONMENT = str(os.getenv("ENVIRONMENT"))

_MOCK_ROUTINE_ID = uuid.UUID("aaaaaaaa-aaaa-aaaa-aaaa-aaaaaaaaaaaa")
_MOCK_HABIT_1_ID = uuid.UUID("bbbbbbbb-bbbb-bbbb-bbbb-bbbbbbbbbbbb")
_MOCK_HABIT_2_ID = uuid.UUID("cccccccc-cccc-cccc-cccc-cccccccccccc")

MOCK_EXPORT_DATA = {
    "routines": [
        SimpleNamespace(
            id=_MOCK_ROUTINE_ID,
            name="Morning Warmup",
            period_of_day="Morning",
            time_of_day="07:00",
        )
    ],
    "habits": [
        SimpleNamespace(
            id=_MOCK_HABIT_1_ID,
            name="Drink Water",
            reminder_time="07:15",
            routine_id=_MOCK_ROUTINE_ID,
        ),
        SimpleNamespace(
            id=_MOCK_HABIT_2_ID,
            name="Journal",
            reminder_time="08:00",
            routine_id=None,
        ),
    ],
    "logs": [
        SimpleNamespace(habit_id=_MOCK_HABIT_1_ID, activity_date=date(2026, 9, 25)),
        SimpleNamespace(habit_id=_MOCK_HABIT_1_ID, activity_date=date(2026, 9, 26)),
        SimpleNamespace(habit_id=_MOCK_HABIT_2_ID, activity_date=date(2026, 9, 26)),
        SimpleNamespace(habit_id=_MOCK_HABIT_1_ID, activity_date=date(2026, 9, 27)),
        SimpleNamespace(habit_id=_MOCK_HABIT_2_ID, activity_date=date(2026, 9, 28)),
    ],
}


class UserService:
    @staticmethod
    def _issue_tokens(user_id, email) -> tuple[str, str]:
        """Create a fresh access + refresh token pair."""
        token_payload = {"sub": str(user_id), "email": email}
        access_token = create_token(data=token_payload, token_type="access")
        refresh_token = create_token(data=token_payload, token_type="refresh")
        return access_token, refresh_token

    @staticmethod
    async def _persist_session(db: AsyncSession, user_id, refresh_token: str) -> None:
        """Decode the refresh token and store a hashed session row."""
        decoded_refresh = decode_token(refresh_token, expected_type="refresh")
        refresh_exp = datetime.fromtimestamp(decoded_refresh["exp"], tz=UTC)

        session_entry = Session(
            user_id=user_id,
            refresh_token_hash=hash_token(refresh_token),
            expires_at=refresh_exp,
        )
        db.add(session_entry)
        await db.commit()

    @staticmethod
    async def _generate_excel_bytes(
        user: User, routines: Routine, habits: Habit, logs: ActivityLog
    ):
        wb = Workbook()

        ws_profile = wb.active
        ws_profile.title = "Profile summary"
        ws_profile.append(["Name", user.full_name])
        ws_profile.append(["Email", user.email])
        ws_profile.append(["Subscription", user.subscription_status])
        ws_profile.append(["Member since", user.created_at.strftime("%Y-%m-%d")])

        ws_routines = wb.create_sheet("Routines")
        ws_routines.append(["Routine name", "Habit", "Remminder"])
        for r in routines:
            for h in habits:
                if h.routine_id == r.id:
                    ws_routines.append([r.name, h.name, str(h.reminder_time or "None")])

        ws_habits = wb.create_sheet("Standalone Habits")
        ws_habits.append(["Habit name", "Reminder"])
        for h in habits:
            if not h.routine_id:
                ws_habits.append([h.name, str(h.reminder_time or "None")])

        ws_logs = wb.create_sheet("Activity History")
        ws_logs.append(["Date", "Habit name"])
        habit_map = {h.id: h.name for h in habits}
        for log in logs:
            ws_logs.append(
                [
                    log.activity_date.strftime("%Y-%m-%d"),
                    habit_map.get(log.habit_id, "Unknown"),
                ]
            )

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @staticmethod
    async def _generate_pdf_bytes(user: User, routines: list[Routine], habits: list[Habit], logs: list[ActivityLog]):
        pdf = FPDF()

        pdf.add_page()
        pdf.set_font("Arial", style="B", size=16)
        pdf.cell(200, 10, text="Rhytm Performance Export", ln=True, align="C")
        pdf.ln(10)

        pdf.set_font("Arial", size=12)
        pdf.cell(200, 10, text=f"Name: {user.full_name}", ln=True)
        pdf.cell(200, 10, text=f"Email: {user.email}", ln=True)
        pdf.cell(200, 10, text=f"Subscription: {user.subscription_status}", ln=True)
        pdf.ln(5)

        pdf.set_font("Arial", style="B", size=14)
        pdf.cell(200, 10, text="Routines:", ln=True)
        pdf.set_font("Arial", size=12)
        for r in routines:
            pdf.cell(200, 10, text=f"  {r.name} ({r.period_of_day}, {r.time_of_day})", ln=True)
            for h in habits:
                if h.routine_id == r.id:
                    pdf.cell(200, 10, text=f"    - {h.name} (Reminder: {h.reminder_time or 'None'})", ln=True)
        pdf.ln(5)

        pdf.set_font("Arial", style="B", size=14)
        pdf.cell(200, 10, text="Standalone Habits:", ln=True)
        pdf.set_font("Arial", size=12)
        for h in habits:
            if not h.routine_id:
                pdf.cell(200, 10, text=f"- {h.name} (Reminder: {h.reminder_time or 'None'})", ln=True)
        pdf.ln(5)

        pdf.set_font("Arial", style="B", size=14)
        pdf.cell(200, 10, text=f"Total Lifetime Check-ins: {len(logs)}", ln=True)
        pdf.set_font("Arial", size=12)
        habit_map = {h.id: h.name for h in habits}
        for log in logs:
            pdf.cell(
                200,
                10,
                text=f"{log.activity_date}: {habit_map.get(log.habit_id, 'Unknown')}",
                ln=True,
            )
        return pdf.output()

    @staticmethod
    async def _process_export_task(
        user_id: str, job_id: str, format: str, meta_key: str, rdb: Redis
    ):
        file_key = f"export_file:{user_id}:{job_id}"
        try:
            async with AsyncSessionLocal() as bg_db:
                user = (
                    await bg_db.execute(select(User).where(User.id == user_id))
                ).scalar_one()

                if ENVIRONMENT == "development":
                    routines = MOCK_EXPORT_DATA["routines"]
                    habits = MOCK_EXPORT_DATA["habits"]
                    logs = MOCK_EXPORT_DATA["logs"]
                else:
                    routines = (
                        await bg_db.execute(
                            select(Routine).where(Routine.user_id == user_id)
                        )
                    ).scalars().all()
                    habits = (
                        await bg_db.execute(select(Habit).where(Habit.user_id == user_id))
                    ).scalars().all()
                    logs = (
                        await bg_db.execute(
                            select(ActivityLog)
                            .join(Habit)
                            .where(Habit.user_id == user_id)
                        )
                    ).scalars().all()

            if format == "excel":
                file_bytes = await UserService._generate_excel_bytes(
                    user, routines, habits, logs
                )
                filename = f"rhytm_export_{datetime.now(UTC).date()}.xlsx"
            else:
                file_bytes = await UserService._generate_pdf_bytes(user, routines, habits, logs)
                filename = f"rhytm_export_{datetime.now(UTC).date()}.pdf"
            await rdb.setex(name=file_key, time=300, value=file_bytes)
            await rdb.setex(
                name=meta_key,
                time=300,
                value=json.dumps(
                    {"status": "ready", "filename": filename, "format": format}
                ),
            )
        except Exception as e:
            logger.exception(f"Failed to generate {format} export for {user_id}")
            await rdb.setex(
                name=meta_key,
                time=300,
                value=json.dumps({"status": "failed", "error": str(e)}),
            )

    async def download_export(
        self, user_id: uuid.UUID, job_id: str, rdb: Redis
    ) -> dict:
        meta_key = f"export_meta:{user_id}:{job_id}"
        file_key = f"export_file:{user_id}:{job_id}"

        raw_meta_data = await rdb.get(meta_key)
        if not raw_meta_data:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=USER_MESSAGES.EXPORT_NOT_FOUND)

        meta_data = json.loads(raw_meta_data)
        if meta_data.get("status") == "failed":
            raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail=USER_MESSAGES.EXPORT_GENERATION_FAILED)

        if meta_data.get("status") == "processing":
            raise HTTPException(status_code=status.HTTP_202_ACCEPTED, detail=USER_MESSAGES.EXPORT_STILL_PROCESSING)

        file_bytes = await rdb.get(file_key)
        if not file_bytes:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=USER_MESSAGES.EXPORT_NOT_FOUND)

        await rdb.delete(meta_key)
        await rdb.delete(file_key)

        content_type = {
            "application/pdf" if meta_data.get("format") == "pdf" 
            else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        }

        return Response(
            content=file_bytes,
            media_type=content_type,
            headers={
                "Content-Disposition": f"attachment; filename={meta_data.get('filename')}"
            },
        )

    async def initiate_signup(
        self, email: str, full_name: str, db: AsyncSession, rdb: Redis
    ) -> dict:
        normalized_email = email.lower()
        result = await db.execute(select(User).where(User.email == normalized_email))
        existing_user = result.scalar_one_or_none()
        if existing_user:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=USER_MESSAGES.USER_EXISTS,
            )

        otp = f"{secrets.randbelow(10000):04d}" if ENVIRONMENT == "prod" else "0000"

        redis_key = f"signup_otp:{normalized_email}"
        signup_data = {
            "full_name": full_name,
            "email": normalized_email,
            "otp": otp,
        }
        await rdb.set(redis_key, json.dumps(signup_data), ex=SIGNUP_OTP_EXPIRY_SECONDS)

        if ENVIRONMENT == "prod":
            send_email_otp(
                to_mail=normalized_email,
                name=full_name,
                otp=otp,
                expiry_minutes=SIGNUP_OTP_EXPIRY_SECONDS // 60,
            )
        else:
            # Log OTP to console for dev environment
            logger.info(f"[SIGNUP OTP] OTP for {normalized_email} is: {otp}")

        message = USER_MESSAGES.OTP_SENT if ENVIRONMENT == "prod" else "Dev otp is 0000"

        return {"message": message}

    async def complete_signup(
        self,
        email: str,
        otp: str,
        password: str,
        db: AsyncSession,
        rdb: Redis,
    ) -> dict:
        normalized_email = email.lower()
        redis_key = f"signup_otp:{normalized_email}"
        cached_data = await rdb.get(redis_key)
        if not cached_data:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=USER_MESSAGES.INVALID_OR_EXPIRED_OTP,
            )

        data = json.loads(cached_data)

        if data.get("otp") != otp.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=USER_MESSAGES.INVALID_OR_EXPIRED_OTP,
            )

        result = await db.execute(select(User).where(User.email == normalized_email))
        if result.scalar_one_or_none():
            await rdb.delete(redis_key)
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=USER_MESSAGES.USER_EXISTS,
            )

        new_user = User(
            full_name=data["full_name"],
            email=normalized_email,
            password=hash_password(password),
        )
        db.add(new_user)
        await db.commit()
        await db.refresh(new_user)

        # Fetch the user and return
        query = await db.execute(
            select(User)
            .options(
                selectinload(User.routines.and_(Routine.deleted_at.is_(None)))
                .selectinload(Routine.habits.and_(Habit.deleted_at.is_(None)))
                .selectinload(Habit.activity_logs),
                selectinload(User.habits.and_(Habit.deleted_at.is_(None))).selectinload(
                    Habit.activity_logs
                ),
            )
            .where(User.email == normalized_email)
        )
        user = query.scalar_one_or_none()

        await rdb.delete(redis_key)

        access_token, refresh_token = self._issue_tokens(user.id, user.email)
        await self._persist_session(db, user.id, refresh_token)

        if ENVIRONMENT == "prod":
            send_welcome_mail(to_mail=normalized_email, name=user.full_name)

        attach_activity(list(user.habits), user.subscription_status)
        for r in user.routines:
            attach_activity(list(r.habits), user.subscription_status)

        return {
            "message": USER_MESSAGES.ACCOUNT_CREATED,
            "access_token": access_token,
            "refresh_token": refresh_token,
            "user": UserResponse.model_validate(user).model_dump(),
        }

    async def login(self, email: str, password: str, db: AsyncSession) -> dict:
        normalized_email = email.lower()

        result = await db.execute(
            select(User)
            .options(
                selectinload(User.routines.and_(Routine.deleted_at.is_(None)))
                .selectinload(Routine.habits.and_(Habit.deleted_at.is_(None)))
                .selectinload(Habit.activity_logs),
                selectinload(User.habits.and_(Habit.deleted_at.is_(None))).selectinload(
                    Habit.activity_logs
                ),
            )
            .where(User.email == normalized_email)
        )
        user = result.scalar_one_or_none()
        if not user or not verify_password(password, user.password):
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail=USER_MESSAGES.INVALID_CREDENTIALS,
            )

        access_token, refresh_token = self._issue_tokens(user.id, user.email)
        await self._persist_session(db, user.id, refresh_token)

        attach_activity(list(user.habits), user.subscription_status)
        for r in user.routines:
            attach_activity(list(r.habits), user.subscription_status)

        return {
            "message": USER_MESSAGES.LOGIN_SUCCESS,
            "access_token": access_token,
            "refresh_token": refresh_token,
            "user": UserResponse.model_validate(user).model_dump(),
        }

    async def request_password_reset(
        self, email: str, db: AsyncSession, rdb: Redis
    ) -> dict:
        normalized_email = email.lower()
        result = await db.execute(select(User).where(User.email == normalized_email))
        user = result.scalar_one_or_none()
        if not user:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=USER_MESSAGES.USER_NOT_FOUND,
            )

        otp = f"{secrets.randbelow(10000):04d}"
        redis_key = f"reset_password_otp:{normalized_email}"
        reset_data = {
            "email": normalized_email,
            "otp": otp,
        }
        await rdb.set(
            redis_key, json.dumps(reset_data), ex=RESET_PASSWORD_OTP_EXPIRY_SECONDS
        )

        # Log OTP to console for now
        # TODO: Set up email OTP sending service (SMTP / SendGrid / Resend / AWS SES)
        logger.info(f"[RESET PASSWORD OTP] OTP for {normalized_email} is: {otp}")

        return {"message": USER_MESSAGES.OTP_SENT}

    async def reset_password(
        self,
        email: str,
        otp: str,
        new_password: str,
        db: AsyncSession,
        rdb: Redis,
    ) -> dict:
        normalized_email = email.lower()
        redis_key = f"reset_password_otp:{normalized_email}"
        cached_data = await rdb.get(redis_key)
        if not cached_data:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=USER_MESSAGES.INVALID_OR_EXPIRED_OTP,
            )

        data = json.loads(cached_data)

        if data.get("otp") != otp.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=USER_MESSAGES.INVALID_OR_EXPIRED_OTP,
            )

        result = await db.execute(select(User).where(User.email == normalized_email))
        user = result.scalar_one_or_none()
        if not user:
            await rdb.delete(redis_key)
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=USER_MESSAGES.USER_NOT_FOUND,
            )

        user.password = hash_password(new_password)
        await db.commit()
        await rdb.delete(redis_key)

        return {"message": USER_MESSAGES.PASSWORD_RESET_SUCCESS}

    async def logout(
        self,
        access_token: str,
        refresh_token: str,
        db: AsyncSession,
        rdb: Redis,
    ) -> dict:
        try:
            access_payload = decode_token(access_token, expected_type="access")
        except jwt.PyJWTError:
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail=USER_MESSAGES.INVALID_TOKEN,
            )

        try:
            refresh_payload = decode_token(refresh_token, expected_type="refresh")
        except jwt.PyJWTError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=USER_MESSAGES.INVALID_TOKEN,
            )

        now = datetime.now(UTC)
        refresh_exp_ts = refresh_payload.get("exp")
        if not refresh_exp_ts or datetime.fromtimestamp(refresh_exp_ts, tz=UTC) <= now:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=USER_MESSAGES.INVALID_TOKEN,
            )

        token_hash = hash_token(refresh_token)
        result = await db.execute(
            select(Session).where(Session.refresh_token_hash == token_hash)
        )
        session_entry = result.scalar_one_or_none()
        if not session_entry:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=USER_MESSAGES.SESSION_EXPIRED,
            )

        if session_entry.expires_at <= now:
            await db.delete(session_entry)
            await db.commit()
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=USER_MESSAGES.SESSION_EXPIRED,
            )

        await db.delete(session_entry)
        await db.commit()

        access_exp = access_payload.get("exp")
        if access_exp and access_exp > now.timestamp():
            ttl = int(access_exp - now.timestamp())
            await rdb.set(f"blacklisted_token:{access_token}", "blacklisted", ex=ttl)

        return {"message": USER_MESSAGES.LOGOUT_SUCCESS}

    async def refresh_tokens(self, refresh_token: str, db: AsyncSession) -> dict:
        try:
            refresh_payload = decode_token(refresh_token, expected_type="refresh")
        except jwt.PyJWTError:
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail=USER_MESSAGES.INVALID_TOKEN,
            )

        now = datetime.now(UTC)
        refresh_exp_ts = refresh_payload.get("exp")
        if not refresh_exp_ts or datetime.fromtimestamp(refresh_exp_ts, tz=UTC) <= now:
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail=USER_MESSAGES.INVALID_TOKEN,
            )

        token_hash = hash_token(refresh_token)
        result = await db.execute(
            select(Session).where(Session.refresh_token_hash == token_hash)
        )
        session_entry = result.scalar_one_or_none()
        if not session_entry:
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail=USER_MESSAGES.SESSION_EXPIRED,
            )

        if session_entry.expires_at <= now:
            await db.delete(session_entry)
            await db.commit()
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail=USER_MESSAGES.SESSION_EXPIRED,
            )

        await db.delete(session_entry)

        user_id = refresh_payload.get("sub")
        email = refresh_payload.get("email")

        new_access_token, new_refresh_token = self._issue_tokens(user_id, email)
        await self._persist_session(db, session_entry.user_id, new_refresh_token)

        return {
            "message": USER_MESSAGES.TOKEN_REFRESH_SUCCESS,
            "access_token": new_access_token,
            "refresh_token": new_refresh_token,
        }

    async def get_profile(self, user_id: str, db: AsyncSession) -> dict:
        result = await db.execute(
            select(User)
            .options(
                selectinload(User.routines.and_(Routine.deleted_at.is_(None)))
                .selectinload(Routine.habits.and_(Habit.deleted_at.is_(None)))
                .selectinload(Habit.activity_logs),
                selectinload(User.habits.and_(Habit.deleted_at.is_(None))).selectinload(
                    Habit.activity_logs
                ),
            )
            .where(User.id == user_id)
        )
        user = result.scalar_one_or_none()
        if not user:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=USER_MESSAGES.USER_NOT_FOUND,
            )
        attach_activity(list(user.habits), user.subscription_status)
        for r in user.routines:
            attach_activity(list(r.habits), user.subscription_status)

        return UserResponse.model_validate(user).model_dump()

    async def delete_account(
        self, user_id: str, access_token: str, db: AsyncSession, rdb: Redis
    ) -> None:
        result = await db.execute(select(User).where(User.id == user_id))
        user = result.scalar_one_or_none()
        if not user:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=USER_MESSAGES.USER_NOT_FOUND,
            )

        await db.execute(delete(Session).where(Session.user_id == user.id))
        await db.delete(user)
        await db.commit()

        try:
            access_payload = decode_token(access_token, expected_type="access")
            access_exp = access_payload.get("exp")
            now = datetime.now(UTC)
            if access_exp and access_exp > now.timestamp():
                ttl = int(access_exp - now.timestamp())
                await rdb.set(
                    f"blacklisted_token:{access_token}", "blacklisted", ex=ttl
                )
        except jwt.PyJWTError:
            pass

    async def generate_export(
        self,
        user_id: str,
        format: str,
        db: AsyncSession,
        rdb: Redis,
    ):
        result = await db.execute(select(User).where(User.id == user_id))
        user = result.scalar_one_or_none()
        if not user:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=USER_MESSAGES.USER_NOT_FOUND,
            )
        if user.subscription_status == "basic":
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN, detail=USER_MESSAGES.USER_NOT_PRO
            )
        if format not in ["pdf", "excel"]:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail=USER_MESSAGES.INVALID_EXPORT_FORMAT,
            )
        job_id = str(uuid.uuid4())
        meta_key = f"export_meta:{user_id}:{job_id}"
        await rdb.setex(
            name=meta_key,
            time=300,
            value=json.dumps({"status": "processing", "format": format}),
        )

        asyncio.create_task(
            UserService._process_export_task(user_id, job_id, format, meta_key, rdb)
        )

        return {"message": "Export generation started", "job_id": job_id, "status": "processing"}
